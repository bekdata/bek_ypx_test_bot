"""
Fayldan savollarni import qilish
Import questions from files (Excel, Word, Text)
"""

import os
from typing import List, Dict


def parse_questions_from_text(text: str) -> List[Dict]:
    """
    Text fayldan savollarni parse qilish
    
    Format:
    1. Savol matni?
    A) Variant A
    B) Variant B
    C) Variant C
    D) Variant D
    Javob: B
    Kategoriya: Kategoriya nomi
    
    2. Keyingi savol...
    """
    questions = []
    lines = text.split('\n')
    current_question = None
    current_options = {}
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Savol raqami va matni
        if line and line[0].isdigit() and '.' in line[:5]:
            # Oldingi savolni saqlash
            if current_question and current_options:
                questions.append({
                    'question': current_question,
                    'options': current_options,
                    'correct_answer': current_correct,
                    'category': current_category
                })
            
            # Yangi savol
            current_question = line.split('.', 1)[1].strip()
            current_options = {}
            current_correct = None
            current_category = "Umumiy"
        
        # Variant
        elif line and line[0] in ['A', 'B', 'C', 'D'] and ')' in line[:3]:
            key = line[0]
            value = line.split(')', 1)[1].strip()
            current_options[key] = value
        
        # To'g'ri javob
        elif line.lower().startswith('javob:'):
            current_correct = line.split(':', 1)[1].strip().upper()
        
        # Kategoriya
        elif line.lower().startswith('kategoriya:'):
            current_category = line.split(':', 1)[1].strip()
        
        i += 1
    
    # Oxirgi savolni qo'shish
    if current_question and current_options:
        questions.append({
            'question': current_question,
            'options': current_options,
            'correct_answer': current_correct,
            'category': current_category
        })
    
    return questions


def parse_questions_from_simple_format(text: str) -> List[Dict]:
    """
    Oddiy formatdan savollarni parse qilish
    Har bir savol bo'sh qator bilan ajratilgan
    """
    questions = []
    blocks = text.strip().split('\n\n')
    
    for block in blocks:
        if not block.strip():
            continue
        
        lines = [l.strip() for l in block.split('\n') if l.strip()]
        
        if len(lines) < 6:  # Kamida: savol + 4 variant + javob
            continue
        
        question_text = lines[0]
        options = {}
        correct_answer = None
        category = "Umumiy"
        
        for line in lines[1:]:
            # Variant
            if line[0] in ['A', 'B', 'C', 'D'] and ')' in line[:3]:
                key = line[0]
                value = line.split(')', 1)[1].strip()
                options[key] = value
            
            # To'g'ri javob
            elif 'javob' in line.lower() or 'to\'g\'ri' in line.lower():
                # Javob: B yoki To'g'ri javob: B
                parts = line.split(':', 1)
                if len(parts) == 2:
                    correct_answer = parts[1].strip().upper()
                    # Faqat birinchi harfni olish
                    if len(correct_answer) > 0:
                        correct_answer = correct_answer[0]
            
            # Kategoriya
            elif 'kategoriya' in line.lower() or 'bo\'lim' in line.lower():
                parts = line.split(':', 1)
                if len(parts) == 2:
                    category = parts[1].strip()
        
        if question_text and len(options) >= 2 and correct_answer:
            questions.append({
                'question': question_text,
                'options': options,
                'correct_answer': correct_answer,
                'category': category
            })
    
    return questions


def validate_questions(questions: List[Dict]) -> tuple[List[Dict], List[str]]:
    """
    Savollarni tekshirish
    Returns: (valid_questions, errors)
    """
    valid_questions = []
    errors = []
    
    for i, q in enumerate(questions, 1):
        error_msgs = []
        
        # Savol matni
        if not q.get('question') or len(q.get('question', '')) < 10:
            error_msgs.append(f"Savol #{i}: Savol matni juda qisqa")
        
        # Variantlar
        options = q.get('options', {})
        if len(options) < 2:
            error_msgs.append(f"Savol #{i}: Kamida 2 ta variant bo'lishi kerak")
        
        # To'g'ri javob
        correct = q.get('correct_answer')
        if not correct or correct not in options:
            error_msgs.append(f"Savol #{i}: To'g'ri javob xato yoki yo'q")
        
        if error_msgs:
            errors.extend(error_msgs)
        else:
            valid_questions.append(q)
    
    return valid_questions, errors


async def process_uploaded_file(file_path: str) -> tuple[List[Dict], List[str]]:
    """
    Yuklangan faylni qayta ishlash
    Returns: (questions, errors)
    """
    file_ext = os.path.splitext(file_path)[1].lower()
    
    try:
        if file_ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # Ikki formatni ham sinab ko'rish
            questions = parse_questions_from_text(content)
            if not questions:
                questions = parse_questions_from_simple_format(content)
            
            return validate_questions(questions)
        
        elif file_ext in ['.doc', '.docx']:
            # Word fayl
            try:
                import docx2txt
                content = docx2txt.process(file_path)
            except ImportError:
                try:
                    from docx import Document
                    doc = Document(file_path)
                    content = '\n'.join([para.text for para in doc.paragraphs])
                except ImportError:
                    return [], ["Python-docx kutubxonasi o'rnatilmagan. pip install python-docx"]
            
            questions = parse_questions_from_simple_format(content)
            return validate_questions(questions)
        
        elif file_ext in ['.xls', '.xlsx']:
            # Excel fayl
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path)
                sheet = wb.active
                
                questions = []
                # Header: Savol | A | B | C | D | Javob | Kategoriya
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if not row[0]:  # Bo'sh qator
                        continue
                    
                    question_text = str(row[0]).strip()
                    options = {}
                    
                    # Variantlar
                    for i, key in enumerate(['A', 'B', 'C', 'D']):
                        if row[i+1]:
                            options[key] = str(row[i+1]).strip()
                    
                    correct_answer = str(row[5]).strip().upper() if row[5] else None
                    category = str(row[6]).strip() if len(row) > 6 and row[6] else "Umumiy"
                    
                    if question_text and options and correct_answer:
                        questions.append({
                            'question': question_text,
                            'options': options,
                            'correct_answer': correct_answer,
                            'category': category
                        })
                
                return validate_questions(questions)
            
            except ImportError:
                return [], ["Openpyxl kutubxonasi o'rnatilmagan. pip install openpyxl"]
        
        else:
            return [], [f"Qo'llab-quvvatlanmaydigan fayl formati: {file_ext}"]
    
    except Exception as e:
        return [], [f"Faylni o'qishda xatolik: {str(e)}"]
